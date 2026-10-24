"""CSV in/out tuned for Excel (UTF-8 with BOM, `;` or `,` auto-detected on read), plus .xlsx
in/out for the Excel buttons (no encoding or locale issues on either platform)."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

Row = dict[str, str]

# Plain decimals only, so the VBA mirror (modMatch.ParseAmount) parses the same values
_AMOUNT_RE = re.compile(r"^[+-]?(?:\d+\.?\d*|\.\d+)$", re.ASCII)


def header_key(name: str) -> str:
    return "_".join(name.strip().lower().replace("-", " ").split())


def _sniff_delimiter(sample: str) -> str:
    first = sample.splitlines()[0] if sample else ""
    counts = {d: first.count(d) for d in (";", ",", "\t")}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] else ","


def read_rows(path: Path) -> list[Row]:
    """Rows keyed by normalised header (lowercase, spaces to underscores)."""
    with path.open(encoding="utf-8-sig", newline="") as f:
        text = f.read()
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text))
    rows = list(reader)
    if not rows:
        return []
    keys = [header_key(h) for h in rows[0]]
    return [
        {k: (r[i] if i < len(r) else "") for i, k in enumerate(keys)}
        for r in rows[1:]
        if any(c.strip() for c in r)
    ]


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, bool):
        return "true" if v else "false"
    if hasattr(v, "isoformat"):
        return str(v.isoformat())
    return str(v)


def read_xlsx_rows(path: Path, sheet: str | None = None) -> list[Row]:
    """First (or named) sheet, header in row 1, keyed like read_rows."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return []
        keys = [header_key(_cell(h)) for h in header]
        return [{k: _cell(v) for k, v in zip(keys, r, strict=False) if k}
                for r in it if any(_cell(v).strip() for v in r)]
    finally:
        wb.close()


def write_xlsx(path: Path, sheets: Mapping[str, tuple[Sequence[str], Iterable[Mapping[str, Any]]]]
               ) -> None:
    """One sheet per entry: header row = field names, values written as text."""
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, (fields, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = name
        first = False
        ws.append(list(fields))
        for r in rows:
            ws.append([_cell(r.get(f, "")) or None for f in fields])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_rows(
    path: Path,
    fieldnames: Sequence[str],
    rows: Iterable[Mapping[str, Any]],
    delimiter: str = ",",
) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(fieldnames), delimiter=delimiter,
                           extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
            n += 1
    return n


def parse_amount(s: str) -> float | None:
    """'1 234,56', '1,234.56', '1234.56' -> 1234.56; blank -> None."""
    t = s.replace(" ", "").replace("\xa0", "").strip()
    if not t:
        return None
    if "," in t and "." in t:
        if t.rfind(",") > t.rfind("."):
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", "")
    else:
        t = t.replace(",", ".")
    return float(t) if _AMOUNT_RE.match(t) else None
