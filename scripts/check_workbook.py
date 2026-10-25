"""Compare what the VBA macros wrote into a saved workbook with the Python result
for the same sheet contents (parity check, any data).

    python scripts/check_workbook.py output/KSeF_IZ_Tool_local.xlsm

Run RunDetect, RunMatch and RunSelfTest in Excel first, then save.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ksefiz import match as m
from ksefiz.cli import DETECT_FIELDS
from ksefiz.io_csv import Row
from ksefiz.iz import detect
from ksefiz.synthetic import DECODED_FIELDS, OPEN_ITEM_FIELDS


def _s(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _block(ws: Worksheet, first_col: int, fields: list[str]) -> list[Row]:
    rows = []
    for r in ws.iter_rows(min_row=2, min_col=first_col, max_col=first_col + len(fields) - 1,
                          values_only=True):
        if any(_s(c).strip() for c in r):
            rows.append({f: _s(c) for f, c in zip(fields, r, strict=True)})
    return rows


def _is_red(cell: Any) -> bool:
    rgb = str(getattr(cell.fill.fgColor, "rgb", "") or "")
    return cell.fill.fill_type == "solid" and rgb.upper().endswith("FFC7CE")


def _first_diff(a: list[tuple[str, ...]], b: list[tuple[str, ...]]) -> str:
    for i, (x, y) in enumerate(zip(a, b, strict=False)):
        if x != y:
            return f"  row {i + 2}\n    vba    {x}\n    python {y}"
    return f"  lengths differ: vba={len(a)} python={len(b)}"


def main(path: Path) -> int:
    wb = load_workbook(path, data_only=True)
    failures = 0

    ws = wb["1_Detect"]
    statements = _block(ws, 1, ["line_id", "date", "amount", "currency", "payer", "free_text"])
    py_det = [
        (s["line_id"], d.check.kind, d.raw.replace("\r", " ").replace("\n", " "),
         d.check.normalised, "Y" if d.check.crc_ok else "N", "Y" if d.check.valid else "N",
         d.check.reason)
        for s in statements for d in detect(s["free_text"])
    ]
    vba_det = [tuple(r.values()) for r in _block(ws, 8, DETECT_FIELDS)]
    if vba_det == py_det:
        print(f"detect: OK, {len(py_det)} rows identical")
    else:
        failures += 1
        print("detect: MISMATCH\n" + _first_diff(vba_det, py_det))

    # Lines without a valid ID must be the ones RunDetect painted red (A:F)
    py_red = {s["line_id"] for s in statements
              if not any(d.check.valid for d in detect(s["free_text"]))}
    vba_red = {_s(row[0].value) for row in ws.iter_rows(min_row=2, max_col=6)
               if row[0].value is not None and _is_red(row[0])}
    if vba_red == py_red:
        print(f"highlight: OK, {len(py_red)} lines without a valid ID are red")
    else:
        failures += 1
        print(f"highlight: MISMATCH red={len(vba_red)} expected={len(py_red)}; "
              f"e.g. missing {sorted(py_red - vba_red)[:3]}, extra {sorted(vba_red - py_red)[:3]}")

    decoded = _block(wb["2_Decode"], 3, DECODED_FIELDS)
    open_items = _block(wb["3_Match"], 1, OPEN_ITEM_FIELDS)
    py_match = [(r.ksef_number, r.status) for r in m.match(decoded, open_items)]
    vba_match = [(r["ksef_number"], r["status"])
                 for r in _block(wb["3_Match"], 8, m.MATCH_FIELDS)]
    if vba_match == py_match:
        print(f"match: OK, {len(py_match)} statuses identical")
    else:
        failures += 1
        print("match: MISMATCH\n" + _first_diff(vba_match, py_match))

    print(f"self-test (_Vectors!K3): {_s(wb['_Vectors']['K3'].value) or '(not run)'}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1])))
