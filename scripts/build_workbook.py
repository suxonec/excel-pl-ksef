"""Build the Excel workbook in one of three flavours.

    python scripts/build_workbook.py           -> excel/KSeF_IZ_Tool_template.xlsx   (committed)
        empty sheets + self-test vectors
    python scripts/build_workbook.py --demo    -> excel/KSeF_IZ_Tool_demo.xlsx       (committed)
        + synthetic statements, decoded rows and open items; nothing machine-specific
    python scripts/build_workbook.py --local   -> output/KSeF_IZ_Tool_local.xlsx     (git-ignored)
        + this machine's helper_path and the KSeF TEST invoices from
        ~/.ksefiz/test/invoices.csv in 4_Generate / 5_Lookup - for your own testing only

Then import excel/vba/*.bas in Excel, run SetupWorkbook and save as .xlsm (see README).
Sheet layouts must match the column constants in excel/vba/*.bas.
"""

from __future__ import annotations

import argparse
import csv
import os
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ksefiz import __version__, operations, synthetic
from ksefiz.cli import DETECT_FIELDS
from ksefiz.match import MATCH_FIELDS

ROOT = Path(__file__).resolve().parent.parent
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
OUTPUT_FILL = PatternFill("solid", fgColor="375623")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FONT = Font(italic=True, color="7F7F7F")

DECODE_FIELDS = operations.DECODE_FIELDS
GENERATE_IN = operations.GENERATE_IN
GENERATE_OUT = operations.GENERATE_FIELDS
LOOKUP_OUT = operations.LOOKUP_FIELDS
QUERY_OUT = operations.QUERY_FIELDS
LOG_FIELDS = ["timestamp", "user", "environment", "operation", "request_count", "input_hash",
              "result", "error_code", "correlation_id"]
HELPER_NOTE = "Calls KSeF through the ksefiz helper (Config > helper_path, environment, profile)."


def _headers(ws: Worksheet, col: int, names: list[str], output: bool = False) -> None:
    for i, name in enumerate(names):
        c = ws.cell(row=1, column=col + i, value=name)
        c.font = HEADER_FONT
        c.fill = OUTPUT_FILL if output else HEADER_FILL
        ws.column_dimensions[c.column_letter].width = max(12, len(name) + 2)


def _text_cols(ws: Worksheet, letters: str) -> None:
    for letter in letters:
        ws.column_dimensions[letter].number_format = "@"


def _width(ws: Worksheet, **widths: int) -> None:
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def _rows(ws: Worksheet, col: int, fields: list[str], rows: list[dict[str, str]]) -> None:
    for r, row in enumerate(rows, start=2):
        for i, f in enumerate(fields):
            # Excel cells hold LF line breaks only; a raw CR would not survive the XML anyway
            v = row.get(f, "").replace("\r\n", "\n").replace("\r", "\n")
            ws.cell(row=r, column=col + i, value=v if v != "" else None)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def build(demo: bool, local: bool = False) -> Workbook:
    wb = Workbook()
    # Fixed metadata: no user name or build time in committed files, reproducible builds
    wb.properties.creator = "ksefiz"
    wb.properties.lastModifiedBy = "ksefiz"
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    readme = wb.active
    assert readme is not None
    readme.title = "README"
    readme["A1"] = '="KSeF IZ Tool  |  ENVIRONMENT: "&Config!B2'
    readme["A1"].font = Font(bold=True, size=16)
    for env, color in (("PROD", "C00000"), ("DEMO", "ED7D31"), ("TEST", "70AD47")):
        readme.conditional_formatting.add(
            "A1:H1", FormulaRule(formula=[f'Config!$B$2="{env}"'],
                                 fill=PatternFill("solid", bgColor=color),
                                 font=Font(color="FFFFFF", bold=True)))
    lines = [
        (3, f"Version {__version__}. Offline modules: M1 Detect, M2 Validate (worksheet "
            "functions), M4 Match. API modules (Decode, Generate, Lookup, Query) come in step 6."
         ),
        (4, "No secrets live in this workbook. The KSeF token stays in the OS keyring / helper."),
        (6, "Buttons (created by macro SetupWorkbook after importing excel/vba/*.bas). API "
            "buttons need Config > helper_path; Mac: run scripts/install_mac_helper.sh once."),
        (15, "How to use"),
        (16, "1. 1_Detect: paste statement lines into A2:F (line_id, date, amount, currency, payer,"
             " free_text). Click Detect IDs. Results land in H:N as values."),
        (17, "2. 2_Decode: paste IZ into A2 or click 'Take IZ from 1_Detect', then Decode. "
             "Results in C:K. Esc cancels a running job."),
        (18, "3. 3_Match: paste the ERP open items extract into A2:F. Click Match. Output: H:S."),
        (19, "4. Worksheet functions: =IZ_VALID(A2), =KSEF_VALID(A2), =IZ_NORMALISE(A2), "
             "=IZ_CRC8(A2), =ID_REASON(A2), =ID_KIND(A2)."),
        (20, "5. Run self-test after any VBA change: it replays the Python test vectors."),
        (21, "6. 4_Generate: invoices in A:F; 'Check (no API)' first, then 'Generate IZ'. "
             "5_Lookup: KSeF numbers in A. 6_Query: dates in B1/B2, own_only in B3."),
        (23, "Match statuses: Matched, Matched (no amount), Amount diff, Currency diff, "
             "Unmatched, Duplicate in ERP. Invoices in several IZ are judged on the total paid."),
    ]
    for row, text in lines:
        readme.cell(row=row, column=1, value=text)
    readme["A15"].font = Font(bold=True)
    readme.column_dimensions["A"].width = 4

    cfg = wb.create_sheet("Config")
    _headers(cfg, 1, ["setting", "value", "notes"])
    settings = [
        ("environment", "TEST", "TEST / DEMO / PROD. PROD also needs --env prod and confirmation"),
        ("context_nip", None, "Synthetic NIP of the TEST context"),
        ("token_source", "keyring", "Never put the token in this workbook"),
        ("helper_path", "ksefiz", "Full path to the ksefiz executable, e.g. .../.venv/bin/ksefiz"),
        ("profile", None, "Stored ksefiz context profile; blank = default (TEST: 'buyer' too)"),
        ("log_folder", None, "Absolute folder for audit.jsonl; blank = ~/.ksefiz/logs"),
        ("max_invoices_per_iz", 500, "Default; effective value from GET /limits/context"),
        ("decode_batch_size", 10, "API maximum per decode call"),
        ("max_iz_per_invoice", 132, "Per context"),
        ("query_max_days", 100, "Period query window"),
        ("budget_left", None, "collectiveIdentifier requests left, updated after each API run"),
    ]
    for r, (k, v, note) in enumerate(settings, start=2):
        cfg.cell(row=r, column=1, value=k)
        cfg.cell(row=r, column=2, value=v)
        cfg.cell(row=r, column=3, value=note).font = NOTE_FONT
    env_list = DataValidation(type="list", formula1='"TEST,DEMO,PROD"', allow_blank=False)
    cfg.add_data_validation(env_list)
    env_list.add("B2")
    _text_cols(cfg, "B")
    _width(cfg, A=22, B=18, C=70)

    det = wb.create_sheet("1_Detect")
    _headers(det, 1, synthetic.STATEMENT_FIELDS)
    _headers(det, 8, DETECT_FIELDS, output=True)
    _text_cols(det, "AFHIJKLMN")
    _width(det, F=80, J=40, K=38)
    det.freeze_panes = "A2"

    dec = wb.create_sheet("2_Decode")
    _headers(dec, 1, ["iz_to_decode"])
    _headers(dec, 3, DECODE_FIELDS, output=True)
    dec["M6"] = HELPER_NOTE
    dec["M6"].font = NOTE_FONT
    _text_cols(dec, "ACD")
    _width(dec, A=38, C=38, D=38, G=30)
    dec.freeze_panes = "A2"

    mat = wb.create_sheet("3_Match")
    _headers(mat, 1, synthetic.OPEN_ITEM_FIELDS)
    _headers(mat, 8, MATCH_FIELDS, output=True)
    _text_cols(mat, "ACHI")
    _width(mat, C=38, H=38, I=38, S=20)
    mat.freeze_panes = "A2"

    gen = wb.create_sheet("4_Generate")
    _headers(gen, 1, GENERATE_IN)
    _headers(gen, 8, GENERATE_OUT, output=True)
    gen["N6"] = HELPER_NOTE
    gen["N6"].font = NOTE_FONT
    _text_cols(gen, "ABI")
    _width(gen, B=38, I=38)

    lk = wb.create_sheet("5_Lookup")
    _headers(lk, 1, ["ksef_number"])
    _headers(lk, 3, LOOKUP_OUT, output=True)
    lk["J4"] = HELPER_NOTE
    lk["J4"].font = NOTE_FONT
    _text_cols(lk, "ACD")
    _width(lk, A=38, C=38, D=38)

    q = wb.create_sheet("6_Query")
    q["A1"], q["A2"], q["A3"] = "date_created_from", "date_created_to", "own_only"
    q["B3"] = False
    q["A4"] = "Ranges over 100 days are split into 100-day windows automatically."
    q["A4"].font = NOTE_FONT
    for i, name in enumerate(QUERY_OUT):
        c = q.cell(row=5, column=1 + i, value=name)
        c.font, c.fill = HEADER_FONT, OUTPUT_FILL
    q["F4"] = HELPER_NOTE
    q["F4"].font = NOTE_FONT
    _width(q, A=38, B=22, C=14, D=26)

    log = wb.create_sheet("Log")
    _headers(log, 1, LOG_FIELDS)
    _width(log, A=20, G=60)
    log.freeze_panes = "A2"

    vec = wb.create_sheet("_Vectors")
    checks = _read_csv(ROOT / "tests" / "vectors.csv")
    detects = _read_csv(ROOT / "tests" / "vectors_detect.csv")
    check_fields = ["input", "kind", "normalised", "crc_ok", "valid", "reason"]
    _headers(vec, 1, check_fields)
    _headers(vec, 8, ["text", "found"])
    _text_cols(vec, "ABCDEFHI")
    _rows(vec, 1, check_fields, checks)
    _rows(vec, 8, ["text", "found"], detects)
    vec["K1"], vec["K2"] = len(checks), len(detects)
    for col in ("A", "H"):
        for cell in vec[col]:
            cell.alignment = Alignment(wrap_text=True)
    vec.sheet_state = "hidden"

    if demo or local:
        ds = synthetic.generate()
        _rows(det, 1, synthetic.STATEMENT_FIELDS, ds.statements)
        _rows(dec, 3, synthetic.DECODED_FIELDS, ds.decoded)
        _rows(mat, 1, synthetic.OPEN_ITEM_FIELDS, ds.open_items)
    if local:  # machine-specific: never commit (output/ is git-ignored)
        helper = ROOT / ".venv" / ("Scripts/ksefiz.exe" if os.name == "nt" else "bin/ksefiz")
        if helper.exists():
            cfg["B5"] = str(helper)
        _local_test_invoices(gen, lk, q)
    return wb


def _local_test_invoices(gen: Worksheet, lk: Worksheet, q: Worksheet) -> None:
    """KSeF TEST invoices from scripts/send_test_invoices.py, if present (TEST data only).
    Uses the small online-sent set, not the 500-invoice batch."""
    path = Path.home() / ".ksefiz" / "test" / "invoices.csv"
    if not path.exists():
        return
    with path.open(encoding="utf-8-sig", newline="") as f:
        inv = [r for r in csv.DictReader(f) if not (r.get("set") or "").startswith("batch")]
    sellers = list(dict.fromkeys(r["seller_nip"] for r in inv))
    own = [r for r in inv if r["seller_nip"] == sellers[0]]
    rows = [{"seller_nip": r["seller_nip"], "ksef_number": r["ksef_number"], "amount": r["gross"],
             "currency": r["currency"], "description": f"Demo payment {r['invoice_number']}",
             "group_key": "DEMO-ALL"} for r in own]
    if len(sellers) > 1:  # one group that must be rejected before the API (71004)
        other = next(r for r in inv if r["seller_nip"] == sellers[1])
        rows += [{"seller_nip": r["seller_nip"], "ksef_number": r["ksef_number"],
                  "group_key": "DEMO-MIXED"} for r in (own[0], other)]
    _rows(gen, 1, GENERATE_IN, rows)
    _rows(lk, 1, ["ksef_number"], [{"ksef_number": r["ksef_number"]} for r in inv])
    today = date.today()
    q["B1"], q["B2"] = today.replace(day=1), today


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--demo", action="store_true", help="generic demo with synthetic data")
    mode.add_argument("--local", action="store_true",
                      help="personal test copy (helper path, TEST invoices) in output/")
    ap.add_argument("--out", type=Path)
    args = ap.parse_args()
    if args.local:
        default = ROOT / "output" / "KSeF_IZ_Tool_local.xlsx"
    elif args.demo:
        default = ROOT / "excel" / "KSeF_IZ_Tool_demo.xlsx"
    else:
        default = ROOT / "excel" / "KSeF_IZ_Tool_template.xlsx"
    out: Path = args.out or default
    out.parent.mkdir(parents=True, exist_ok=True)
    build(args.demo, args.local).save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
